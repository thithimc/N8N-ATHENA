import json,collections,unicodedata,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import os
# Usage : DEPT=77 CENTRE_NOM="Champs-sur-Marne" python3 prospection/export_excel.py
DEPT=os.environ.get('DEPT','94'); CENTRE_NOM=os.environ.get('CENTRE_NOM')
R=json.load(open(f'prospection/result{DEPT}.json'))
AVEC_DIST=any(r.get('Distance_km')!='' for r in R)
n=lambda s:unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()
forms=collections.defaultdict(collections.Counter)
for r in R: forms[n(r['Ville'])][r['Ville']]+=1
best={k:max(c,key=lambda f:(sum(ch>'\x7f' for ch in f),c[f])) for k,c in forms.items()}
for r in R: r['Ville']=best[n(r['Ville'])]
# dédoublonnage offres / entreprises : on garde tout, mais tri
if AVEC_DIST: R.sort(key=lambda r:(r['Distance_km'] if r['Distance_km']!='' else 999,r['Entreprise'].lower()))
else: R.sort(key=lambda r:(r['Catégorie'],r['Ville'],r['Entreprise'].lower()))
cols=(['Distance_km'] if AVEC_DIST else [])+['Entreprise','Catégorie','Filière','Ville','CP','Adresse','Secteur','Métiers','Type','Intitulé','Niveau','Contrat','Publication','Organisme_école','SIRET','Pertinence','Lien']
W={'Distance_km':11,'Entreprise':34,'Catégorie':30,'Filière':12,'Ville':22,'CP':8,'Adresse':38,'Secteur':40,'Métiers':60,'Type':30,'Intitulé':40,'Niveau':18,'Contrat':22,'Publication':12,'Organisme_école':10,'SIRET':16,'Pertinence':10,'Lien':50}
wb=Workbook(); first=True
def sheet(name,rows,c=cols):
    global first
    ws=wb.active if first else wb.create_sheet(); first=False; ws.title=name
    ws.append(c)
    for x in rows: ws.append([x.get(k,'') for k in c])
    for i,k in enumerate(c,1):
        ws.column_dimensions[get_column_letter(i)].width=W.get(k,15)
        ws.cell(1,i).font=Font(bold=True,color='FFFFFF'); ws.cell(1,i).fill=PatternFill('solid',fgColor='1F4E78')
    if 'Lien' in c:
        li=c.index('Lien')+1
        for row in range(2,ws.max_row+1):
            cell=ws.cell(row,li); cell.hyperlink=cell.value; cell.font=Font(color='0563C1',underline='single')
    ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
offres=[r for r in R if not r['Type'].startswith('Entreprise')]
if not AVEC_DIST: offres.sort(key=lambda r:r['Publication'],reverse=True)
ent=[r for r in R if r['Type'].startswith('Entreprise')]
c2=[k for k in cols if k not in ('Type','Intitulé','Niveau','Publication','Organisme_école')]
sheet('Offres en cours',offres,[k for k in cols if k!='SIRET' and k!='Pertinence'])
sheet('Entreprises MCO',[r for r in ent if 'MCO' in r['Filière']],c2)
sheet('Entreprises NDRC',[r for r in ent if 'NDRC' in r['Filière']],c2)
sheet('Tout',R)
if AVEC_DIST:
    sheet(f'< 15 km {CENTRE_NOM or "du centre"}'[:31],[r for r in R if r['Distance_km']!='' and r['Distance_km']<=15])
ws=wb.create_sheet('Par ville'); ws.append(['Ville','Nb total','MCO','NDRC','Offres en cours'])
agg=collections.defaultdict(lambda:[0,0,0,0])
for r in R:
    a=agg[r['Ville']]; a[0]+=1; a[1]+='MCO' in r['Filière']; a[2]+='NDRC' in r['Filière']; a[3]+=not r['Type'].startswith('Entreprise')
for v,a in sorted(agg.items(),key=lambda x:-x[1][0]): ws.append([v]+a)
for i in range(1,6): ws.cell(1,i).font=Font(bold=True)
ws.column_dimensions['A'].width=28
ws=wb.create_sheet('Par catégorie'); ws.append(['Catégorie','Nb total','MCO','NDRC','Offres en cours','Exemples d\'entreprises'])
aggc=collections.defaultdict(lambda:[0,0,0,0,collections.Counter()])
for r in R:
    a=aggc[r['Catégorie']]; a[0]+=1; a[1]+='MCO' in r['Filière']; a[2]+='NDRC' in r['Filière']; a[3]+=not r['Type'].startswith('Entreprise'); a[4][r['Entreprise']]+=1
for v,a in sorted(aggc.items(),key=lambda x:-x[1][0]): ws.append([v]+a[:4]+[', '.join(e for e,_ in a[4].most_common(8))])
for i in range(1,7): ws.cell(1,i).font=Font(bold=True)
ws.column_dimensions['A'].width=45; ws.column_dimensions['F'].width=110
wb.move_sheet('Par catégorie',offset=-(len(wb.sheetnames)-1))
if AVEC_DIST:
    near=[n for n in wb.sheetnames if n.startswith('< 15 km')][0]
    wb.move_sheet(near,offset=1-wb.sheetnames.index(near))
wb.save(f'prospection/entreprises_alternance_BTS_MCO_NDRC_{DEPT}.xlsx')
import csv
with open(f'prospection/entreprises_alternance_BTS_MCO_NDRC_{DEPT}.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=cols,delimiter=';',extrasaction='ignore'); w.writeheader(); w.writerows(R)
print(sorted(((a[0],v) for v,a in aggc.items()),reverse=True));print(len(R),len(offres),len(ent),sum('MCO' in r['Filière'] for r in ent),sum('NDRC' in r['Filière'] for r in ent),len(agg))
for r in offres[:12]: print(r['Publication'],r['Entreprise'],'|',r['Intitulé'],'|',r['Ville'],'|',r['Niveau'])
print(collections.Counter(r['Entreprise'] for r in ent).most_common(25))
